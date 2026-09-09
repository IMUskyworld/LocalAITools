#!/usr/bin/env python3
"""Generate a PPTX file from a JSON spec.

Usage: python make_ppt.py <spec.json>
Spec format:
{
  "path": "D:/out.pptx",
  "title": "演示文稿标题",
  "subtitle": "副标题",
  "slides": [
    {"title": "页面标题", "bullets": ["要点1", "要点2", "要点3"]},
    ...
  ]
}
"""
import json
import sys
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x33, 0x33, 0x33)


def build(spec):
    path = spec["path"]
    title = spec.get("title", "演示文稿")
    subtitle = spec.get("subtitle", "")
    slides = spec.get("slides", [])

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # --- Title slide ---
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    t = slide.shapes.title
    t.text = title
    t.text_frame.paragraphs[0].font.size = Pt(44)
    t.text_frame.paragraphs[0].font.bold = True
    t.text_frame.paragraphs[0].font.color.rgb = ACCENT
    if subtitle:
        st = slide.placeholders[1]
        st.text = subtitle
        st.text_frame.paragraphs[0].font.size = Pt(24)
        st.text_frame.paragraphs[0].font.color.rgb = DARK

    # --- Content slides ---
    for s in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        st = slide.shapes.title
        st.text = s.get("title", "")
        st.text_frame.paragraphs[0].font.size = Pt(32)
        st.text_frame.paragraphs[0].font.bold = True
        st.text_frame.paragraphs[0].font.color.rgb = ACCENT

        body = slide.placeholders[1].text_frame
        body.clear()
        bullets = s.get("bullets", [])
        for i, b in enumerate(bullets):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = b
            p.font.size = Pt(20)
            p.font.color.rgb = DARK
            p.level = 0

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    prs.save(path)
    return {"success": True, "path": path, "slides": len(slides)}


if __name__ == "__main__":
    spec_path = sys.argv[1]
    with open(spec_path, "r", encoding="utf-8") as f:
        spec = json.load(f)
    result = build(spec)
    print(json.dumps(result, ensure_ascii=False))
