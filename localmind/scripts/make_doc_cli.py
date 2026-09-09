#!/usr/bin/env python3
"""LocalMind 文档生成统一 CLI 入口（PyInstaller 打包用，自包含无外部 import）。

用法: make_doc.exe --type ppt|docx|xlsx|pdf <spec.json>

合并 4 个 make_*.py 的 build() 逻辑（自包含，不依赖同目录其他模块，
便于 PyInstaller 单文件打包）。依赖: python-pptx / python-docx / openpyxl / fpdf2
"""
import argparse
import json
import os
import sys


# ========== PPT (from make_ppt.py) ==========

def build_ppt(spec):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    ACCENT = RGBColor(0x1F, 0x4E, 0x79)
    DARK = RGBColor(0x33, 0x33, 0x33)

    path = spec["path"]
    title = spec.get("title", "演示文稿")
    subtitle = spec.get("subtitle", "")
    slides = spec.get("slides", [])

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # --- Title slide ---
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    t = slide.shapes.title
    t.text = title
    t.text_frame.paragraphs[0].font.size = Pt(44)
    t.text_frame.paragraphs[0].font.bold = True
    t.text_frame.paragraphs[0].font.color.rgb = ACCENT
    if subtitle:
        st = slide.placeholders[1]
        st.text = subtitle
        st.text_frame.paragraphs[0].font.size = Pt(24)
        st.text_frame.paragraphs[0].font.color.rgb = DARK

    # --- Content slides ---
    for s in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        st = slide.shapes.title
        st.text = s.get("title", "")
        st.text_frame.paragraphs[0].font.size = Pt(32)
        st.text_frame.paragraphs[0].font.bold = True
        st.text_frame.paragraphs[0].font.color.rgb = ACCENT

        body = slide.placeholders[1].text_frame
        body.clear()
        bullets = s.get("bullets", [])
        for i, b in enumerate(bullets):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = b
            p.font.size = Pt(20)
            p.font.color.rgb = DARK
            p.level = 0

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    prs.save(path)
    return {"success": True, "path": path, "slides": len(slides)}


# ========== DOCX (from make_docx.py) ==========

def build_docx(spec):
    from docx import Document

    path = spec["path"]
    doc = Document()

    title = spec.get("title", "")
    if title:
        doc.add_heading(title, 0)

    meta = spec.get("meta", "")
    if meta:
        doc.add_paragraph(meta).italic = True

    for h in spec.get("headings", []):
        level = h.get("level", 1)
        doc.add_heading(h.get("text", ""), level=level)

    for p in spec.get("paragraphs", []):
        doc.add_paragraph(p)

    for b in spec.get("bullets", []):
        doc.add_paragraph(b, style="List Bullet")

    for n in spec.get("numbered", []):
        doc.add_paragraph(n, style="List Number")

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    doc.save(path)
    return {"success": True, "path": path, "paragraphs": len(spec.get("paragraphs", []))}


# ========== XLSX (from make_xlsx.py) ==========

def build_xlsx(spec):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    path = spec["path"]
    wb = Workbook()
    ws = wb.active
    ws.title = spec.get("sheet_name", "Sheet1")

    headers = spec.get("headers", [])
    rows = spec.get("rows", [])

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill

    for r, row in enumerate(rows, start=2):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)

    for col in range(1, len(headers) + 1):
        max_len = len(str(headers[col - 1])) if headers else 0
        for r in range(2, len(rows) + 2):
            v = ws.cell(row=r, column=col).value
            if v is not None:
                max_len = max(max_len, len(str(v)))
        ws.column_dimensions[chr(64 + col) if col <= 26 else "A"].width = max_len + 4

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)
    return {"success": True, "path": path, "rows": len(rows)}


# ========== PDF (from make_pdf.py) ==========

def _find_cjk_font():
    """查找支持中文的字体（Windows 系统字体）"""
    candidates = [
        r"C:\Windows\Fonts\msyh.ttc",    # 微软雅黑
        r"C:\Windows\Fonts\simhei.ttf",   # 黑体
        r"C:\Windows\Fonts\simsun.ttc",   # 宋体
        r"C:\Windows\Fonts\Deng.ttf",     # 等线
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    return ""


def build_pdf(spec):
    from fpdf import FPDF

    path = spec["path"]
    pdf = FPDF()
    pdf.add_page()

    font_path = _find_cjk_font()
    use_cjk = bool(font_path)
    if use_cjk:
        pdf.add_font("CJK", "", font_path)

    title = spec.get("title", "")
    if title:
        pdf.set_font("CJK" if use_cjk else "helvetica", "", 22)
        pdf.cell(0, 16, title, ln=True, align="C")
        pdf.ln(4)

    subtitle = spec.get("subtitle", "")
    if subtitle:
        pdf.set_font("CJK" if use_cjk else "helvetica", "", 13)
        pdf.cell(0, 10, subtitle, ln=True, align="C")
        pdf.ln(8)

    for section in spec.get("sections", []):
        heading = section.get("heading", "")
        if heading:
            pdf.set_font("CJK" if use_cjk else "helvetica", "", 15)
            pdf.cell(0, 12, heading, ln=True)
            pdf.ln(2)
        pdf.set_font("CJK" if use_cjk else "helvetica", "", 11)
        for body in section.get("body", []):
            pdf.multi_cell(0, 7, body)
            pdf.ln(2)
        pdf.ln(4)

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    pdf.output(path)
    return {"success": True, "path": path, "sections": len(spec.get("sections", []))}


# ========== 分发 ==========

BUILDERS = {
    "ppt": build_ppt,
    "pptx": build_ppt,
    "docx": build_docx,
    "doc": build_docx,
    "word": build_docx,
    "xlsx": build_xlsx,
    "xls": build_xlsx,
    "excel": build_xlsx,
    "pdf": build_pdf,
}


def main():
    parser = argparse.ArgumentParser(description="LocalMind 文档生成器")
    parser.add_argument("--type", required=True, help="文档类型: ppt/docx/xlsx/pdf")
    parser.add_argument("spec", help="spec JSON 文件路径")
    args = parser.parse_args()

    builder = BUILDERS.get(args.type.lower())
    if builder is None:
        print(json.dumps({
            "success": False,
            "error": f"不支持的文档类型: {args.type}（支持 ppt/docx/xlsx/pdf）",
        }, ensure_ascii=False))
        sys.exit(1)

    try:
        with open(args.spec, "r", encoding="utf-8") as f:
            spec = json.load(f)
    except json.JSONDecodeError as e:
        # Debug: dump the file content so we can see what Rust actually wrote
        try:
            with open(args.spec, "r", encoding="utf-8") as f:
                content = f.read()
        except Exception:
            content = "<unreadable>"
        print(json.dumps({
            "success": False,
            "error": f"spec JSON 解析失败: {e}",
            "file_content_preview": content[:300],
        }, ensure_ascii=False))
        sys.exit(1)

    result = builder(spec)
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
