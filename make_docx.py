#!/usr/bin/env python3
"""Generate a Word (.docx) file from a JSON spec.

Usage: python make_docx.py <spec.json>
Spec format:
{
  "path": "D:/out.docx",
  "title": "文档标题",
  "paragraphs": ["段落1", "段落2", ...],
  "headings": [{"level": 1, "text": "章节标题"}, ...],
  "bullets": ["要点1", "要点2", ...]
}
"""
import json
import sys
import os

from docx import Document


def build(spec):
    path = spec["path"]
    doc = Document()

    # Title
    title = spec.get("title", "")
    if title:
        doc.add_heading(title, 0)

    # Subtitle / meta
    meta = spec.get("meta", "")
    if meta:
        doc.add_paragraph(meta).italic = True

    # Headings (in order)
    for h in spec.get("headings", []):
        level = h.get("level", 1)
        doc.add_heading(h.get("text", ""), level=level)

    # Paragraphs
    for p in spec.get("paragraphs", []):
        doc.add_paragraph(p)

    # Bullets
    for b in spec.get("bullets", []):
        doc.add_paragraph(b, style="List Bullet")

    # Numbered list
    for n in spec.get("numbered", []):
        doc.add_paragraph(n, style="List Number")

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    doc.save(path)
    return {"success": True, "path": path, "paragraphs": len(spec.get("paragraphs", []))}


if __name__ == "__main__":
    spec_path = sys.argv[1]
    with open(spec_path, "r", encoding="utf-8") as f:
        spec = json.load(f)
    result = build(spec)
    print(json.dumps(result, ensure_ascii=False))
